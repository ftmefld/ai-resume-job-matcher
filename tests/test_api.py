from io import BytesIO

from docx import Document

import app.main as main_module

from app.services.matcher_pipeline import (
    MatcherUnavailableError,
)

from app.services.matching_models import (
    FinalMatchResult,
    RequirementMatch,
)


FAKE_ANALYSIS = FinalMatchResult(
    match_score=50.0,

    score_breakdown={
        "skill": 100.0,
        "tool": 0.0,
    },

    requirement_matches=[
        RequirementMatch(
            requirement_id="R001",
            requirement="Python",
            category="skill",
            importance="important",
            status="matched",
            evidence=[
                "Developed applications using Python."
            ],
            rationale=(
                "Python experience is explicitly "
                "supported by the resume."
            ),
        ),

        RequirementMatch(
            requirement_id="R002",
            requirement="Docker",
            category="tool",
            importance="important",
            status="missing",
            evidence=[],
            rationale=(
                "Docker experience is not supported "
                "by the resume."
            ),
        ),
    ],

    matched_requirements=[
        "Python"
    ],

    partial_requirements=[],

    missing_requirements=[
        "Docker"
    ],

    suggestions=[
        "Gain hands-on Docker experience."
    ],
)


def fake_analysis(
    resume_text,
    job_description,
):
    return FAKE_ANALYSIS


def create_test_docx():
    document = Document()

    document.add_paragraph(
        "Developed applications using Python."
    )

    document.add_paragraph(
        "Experience with machine learning and APIs."
    )

    file_stream = BytesIO()

    document.save(
        file_stream
    )

    file_stream.seek(0)

    return file_stream.getvalue()


def test_health_endpoint(
    client,
):
    response = client.get(
        "/health"
    )

    assert response.status_code == 200

    data = response.json()

    assert data["status"] == "healthy"
    assert data["version"] == "1.0.0"
    assert data["environment"] == "development"


def test_analyze_text_success(
    client,
    monkeypatch,
):
    monkeypatch.setattr(
        main_module,
        "analyze_resume_v2",
        fake_analysis,
    )

    response = client.post(
        "/analyze",
        json={
            "resume_text": (
                "Developed applications using Python."
            ),
            "job_description": (
                "Python and Docker experience required."
            ),
        },
    )

    assert response.status_code == 200

    data = response.json()

    assert data["analysis_id"] is not None

    assert data["match_score"] == 50.0

    assert data["matched_requirements"] == [
        "Python"
    ]

    assert data["partial_requirements"] == []

    assert data["missing_requirements"] == [
        "Docker"
    ]

    assert (
        data["score_breakdown"]["skill"]
        == 100.0
    )

    assert (
        data["score_breakdown"]["tool"]
        == 0.0
    )

    assert len(
        data["requirement_matches"]
    ) == 2

    first_requirement = (
        data["requirement_matches"][0]
    )

    assert (
        first_requirement["requirement_id"]
        == "R001"
    )

    assert (
        first_requirement["status"]
        == "matched"
    )

    assert first_requirement["evidence"] == [
        "Developed applications using Python."
    ]


def test_analysis_saved_to_database(
    client,
    monkeypatch,
):
    monkeypatch.setattr(
        main_module,
        "analyze_resume_v2",
        fake_analysis,
    )

    response = client.post(
        "/analyze",
        json={
            "resume_text": (
                "Developed applications using Python."
            ),
            "job_description": (
                "Python and Docker experience required."
            ),
        },
    )

    assert response.status_code == 200

    analysis_id = response.json()[
        "analysis_id"
    ]

    history_response = client.get(
        "/analyses"
    )

    assert history_response.status_code == 200

    history = history_response.json()

    assert len(history) == 1

    saved_analysis = history[0]

    assert (
        saved_analysis["analysis_id"]
        == analysis_id
    )

    assert (
        saved_analysis["match_score"]
        == 50.0
    )

    assert (
        saved_analysis["input_type"]
        == "text"
    )

    assert (
        saved_analysis["resume_filename"]
        is None
    )

    assert (
        saved_analysis["resume_file_type"]
        is None
    )

    assert (
        saved_analysis[
            "matched_requirements"
        ]
        == ["Python"]
    )

    assert (
        saved_analysis[
            "missing_requirements"
        ]
        == ["Docker"]
    )

    assert (
        saved_analysis[
            "partial_requirements"
        ]
        == []
    )

    assert (
        saved_analysis[
            "score_breakdown"
        ]["skill"]
        == 100.0
    )

    assert len(
        saved_analysis[
            "requirement_matches"
        ]
    ) == 2


def test_empty_resume_rejected(
    client,
):
    response = client.post(
        "/analyze",
        json={
            "resume_text": "   ",
            "job_description": (
                "Python developer required."
            ),
        },
    )

    assert response.status_code == 400


def test_invalid_resume_file_rejected(
    client,
):
    files = {
        "resume_file": (
            "resume.txt",
            b"This is a resume.",
            "text/plain",
        )
    }

    data = {
        "job_description": (
            "Python developer required."
        )
    }

    response = client.post(
        "/analyze-file",
        files=files,
        data=data,
    )

    assert response.status_code == 400


def test_docx_upload_success(
    client,
    monkeypatch,
):
    monkeypatch.setattr(
        main_module,
        "analyze_resume_v2",
        fake_analysis,
    )

    docx_bytes = create_test_docx()

    files = {
        "resume_file": (
            "test_resume.docx",
            docx_bytes,
            (
                "application/vnd.openxmlformats-"
                "officedocument."
                "wordprocessingml.document"
            ),
        )
    }

    data = {
        "job_description": (
            "Python and Docker experience required."
        )
    }

    response = client.post(
        "/analyze-file",
        files=files,
        data=data,
    )

    assert response.status_code == 200

    response_data = response.json()

    assert (
        response_data["match_score"]
        == 50.0
    )

    assert (
        response_data[
            "matched_requirements"
        ]
        == ["Python"]
    )

    assert (
        response_data[
            "missing_requirements"
        ]
        == ["Docker"]
    )

    analysis_id = response_data[
        "analysis_id"
    ]

    history_response = client.get(
        f"/analyses/{analysis_id}"
    )

    assert history_response.status_code == 200

    saved_analysis = (
        history_response.json()
    )

    assert (
        saved_analysis["input_type"]
        == "file"
    )

    assert (
        saved_analysis["resume_filename"]
        == "test_resume.docx"
    )

    assert (
        saved_analysis["resume_file_type"]
        == "docx"
    )

    assert (
        saved_analysis[
            "partial_requirements"
        ]
        == []
    )

    assert len(
        saved_analysis[
            "requirement_matches"
        ]
    ) == 2


def test_llm_failure_returns_503(
    client,
    monkeypatch,
):
    def fake_failure(
        resume_text,
        job_description,
    ):
        raise MatcherUnavailableError(
            "The AI analysis service is "
            "temporarily unavailable."
        )

    monkeypatch.setattr(
        main_module,
        "analyze_resume_v2",
        fake_failure,
    )

    response = client.post(
        "/analyze",
        json={
            "resume_text": (
                "Experienced Python developer."
            ),
            "job_description": (
                "Python developer required."
            ),
        },
    )

    assert response.status_code == 503

    assert (
        response.json()["detail"]
        == (
            "The AI analysis service is "
            "temporarily unavailable."
        )
    )


def test_llm_failure_not_saved(
    client,
    monkeypatch,
):
    def fake_failure(
        resume_text,
        job_description,
    ):
        raise MatcherUnavailableError(
            "The AI analysis service is "
            "temporarily unavailable."
        )

    monkeypatch.setattr(
        main_module,
        "analyze_resume_v2",
        fake_failure,
    )

    response = client.post(
        "/analyze",
        json={
            "resume_text": (
                "Experienced Python developer."
            ),
            "job_description": (
                "Python developer required."
            ),
        },
    )

    assert response.status_code == 503

    history_response = client.get(
        "/analyses"
    )

    assert history_response.status_code == 200

    assert history_response.json() == []